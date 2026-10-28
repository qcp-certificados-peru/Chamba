"""Read the browser's real XLSX bytes with an independent spreadsheet library."""

from io import BytesIO
import json
from pathlib import Path
import shutil
import subprocess
from zipfile import ZipFile

from openpyxl import load_workbook
import pytest


EXPORTER = Path(__file__).resolve().parents[1] / "app/static/links-excel.js"
RUN_EXPORTER = """
const fs = require('node:fs'), vm = require('node:vm');
const context = { window: {}, TextEncoder, Uint8Array, DataView, Blob };
vm.runInNewContext(fs.readFileSync(process.argv[1], 'utf8'), context);
const blob = context.window.QCPLinksExcel.create(JSON.parse(process.argv[2]));
blob.arrayBuffer().then(data => process.stdout.write(Buffer.from(data)));
"""


def workbook_bytes(rows):
    node = shutil.which("node")
    assert node, "Las pruebas del Excel de enlaces necesitan Node.js 18 o superior."
    return subprocess.run([node, "-e", RUN_EXPORTER, str(EXPORTER), json.dumps(rows)],
                          check=True, capture_output=True, timeout=10).stdout


def drive_url(index):
    return f"https://drive.google.com/file/d/synthetic_file_{index:04d}/view"


def test_ten_certificates_export_two_columns_in_order_with_clickable_drive_links():
    rows = [{"filename": f"QCP-{index}-2026 Certificado Perú.pdf", "driveUrl": drive_url(index)} for index in range(10)]
    data = workbook_bytes(rows)
    with ZipFile(BytesIO(data)) as archive:
        assert archive.testzip() is None
        assert not any(name.endswith("vbaProject.bin") for name in archive.namelist())
    book = load_workbook(BytesIO(data))
    assert book.sheetnames == ["Certificados"]
    sheet = book.active
    assert sheet.max_row == 11 and sheet.max_column == 2
    assert list(sheet.values) == [("Nombre del PDF", "Enlace de Google Drive")] + [(row["filename"], row["driveUrl"]) for row in rows]
    assert sheet.freeze_panes == "A2" and sheet.auto_filter.ref == "A1:B11"
    for index, row in enumerate(rows, start=2):
        assert sheet[f"B{index}"].hyperlink.target == row["driveUrl"]
        assert sheet[f"A{index}"].data_type == "s" and sheet[f"B{index}"].data_type == "s"


def test_names_preserve_unicode_and_xml_characters_without_running_formulas():
    names = ['=HYPERLINK("https://example.com","X").pdf', '+SUM(1,2).pdf',
             '-10.pdf', '@nombre.pdf', 'Presión & temperatura <1> "Perú" 😀.pdf']
    data = workbook_bytes([{"filename": name, "driveUrl": drive_url(index)} for index, name in enumerate(names)])
    sheet = load_workbook(BytesIO(data)).active
    for index, name in enumerate(names, start=2):
        assert sheet[f"A{index}"].value == name
        assert sheet[f"A{index}"].data_type == "s"
    with ZipFile(BytesIO(data)) as archive:
        assert b"<f>" not in archive.read("xl/worksheets/sheet1.xml")


@pytest.mark.parametrize("rows", [[], [{"filename": "test.pdf", "driveUrl": "https://portal.example.com/q/test"}],
                                    [{"filename": "test.pdf", "driveUrl": "javascript:alert(1)"}]])
def test_export_rejects_empty_or_non_drive_links(rows):
    with pytest.raises(subprocess.CalledProcessError):
        workbook_bytes(rows)
